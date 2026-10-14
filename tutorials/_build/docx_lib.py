"""Tiny A4 layout toolkit on top of python-docx for the study tutorials."""
import os, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

RASP = "C51A4A"; NAVY = "1F3A5F"; INK = "1F2933"; GREY = "6B7785"; GREEN = "2E8B57"; AMBER = "D98E04"; PURPLE = "7A4FA0"
BODY_FONT = "Calibri"; MONO = "Courier New"
CONTENT_W = 17.2  # cm


def rgb(h): return RGBColor.from_string(h)

def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")): tcPr.remove(e)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)

def borders(cell, **kw):
    """kw: top/left/bottom/right = (size_eighths_pt, hexcolor) or None for none"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None: b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        spec = kw.get(side, None)
        el = OxmlElement(f"w:{side}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:space"), "0"); el.set(qn("w:color"), spec[1])
        b.append(el)

def cell_margins(cell, top=60, bottom=60, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def no_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); trPr.append(e)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); trPr.append(e)

def set_col_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] / 2.54 * 1440)))

def para_fmt(p, before=0, after=4, line=1.12, keep_next=False, keep_lines=False, align=None):
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = line
    f.keep_with_next = keep_next; f.keep_together = keep_lines
    if align is not None: p.alignment = align

def add_runs(p, text, size=None, color=None, bold=None, italic=None, font=None):
    """Mini-markup: **bold**, `code`, __italic__ inside text."""
    import re
    toks = re.split(r"(\*\*.+?\*\*|`.+?`|__.+?__)", text)
    for t in toks:
        if not t: continue
        if t.startswith("**"): r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`"):
            r = p.add_run(t[1:-1]); r.font.name = MONO; r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            r.font.size = Pt((size or 10.5) - 0.7); r.font.color.rgb = rgb("9B1C45"); continue
        elif t.startswith("__"): r = p.add_run(t[2:-2]); r.italic = True
        else: r = p.add_run(t)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = rgb(color)
        if bold is not None and not t.startswith("**"): r.bold = bold
        if italic is not None and not t.startswith("__"): r.italic = italic
        if font: r.font.name = font
    return p


class Doc:
    def __init__(self, header_text, footer_text):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width = Cm(21.0); s.page_height = Cm(29.7)
        s.left_margin = s.right_margin = Cm(1.9); s.top_margin = Cm(1.9); s.bottom_margin = Cm(1.7)
        s.header_distance = Cm(0.8); s.footer_distance = Cm(0.7)
        st = self.d.styles["Normal"]; st.font.name = BODY_FONT; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT); st.font.color.rgb = rgb(INK)
        for name, size, color, before, after in (("Heading 1", 16, NAVY, 14, 6), ("Heading 2", 12.5, RASP, 10, 4), ("Heading 3", 11, NAVY, 8, 3)):
            h = self.d.styles[name]; h.font.name = BODY_FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = rgb(color)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT); h.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT); h.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True
        self.fig_n = 0
        self._header_footer(header_text, footer_text)

    # ---------------- header / footer
    def _field(self, run, code):
        for t, txt in (("begin", None), (None, code), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); run._r.append(e)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)

    def _header_footer(self, header_text, footer_text):
        s = self.d.sections[0]
        hp = s.header.paragraphs[0]; hp.text = ""
        r = hp.add_run(header_text); r.font.size = Pt(8.5); r.font.color.rgb = rgb(GREY)
        pPr = hp._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "3"), ("color", "C9D1D9")): bt.set(qn(f"w:{k}"), v)
        pb.append(bt); pPr.append(pb)
        fp = s.footer.paragraphs[0]; fp.text = ""
        r = fp.add_run(footer_text + "    |    Page "); r.font.size = Pt(8.5); r.font.color.rgb = rgb(GREY)
        r2 = fp.add_run(); r2.font.size = Pt(8.5); r2.font.color.rgb = rgb(GREY); self._field(r2, "PAGE")
        r3 = fp.add_run(" of "); r3.font.size = Pt(8.5); r3.font.color.rgb = rgb(GREY)
        r4 = fp.add_run(); r4.font.size = Pt(8.5); r4.font.color.rgb = rgb(GREY); self._field(r4, "NUMPAGES")
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- text
    def h1(self, t): return self.d.add_heading(t, level=1)
    def h2(self, t): return self.d.add_heading(t, level=2)
    def h3(self, t): return self.d.add_heading(t, level=3)

    def p(self, text, size=None, color=None, italic=None, after=4, align=None, keep_next=False, bold=None):
        p = self.d.add_paragraph(); add_runs(p, text, size=size, color=color, italic=italic, bold=bold)
        para_fmt(p, after=after, align=align, keep_next=keep_next); return p

    def bullets(self, items, size=None, marker="•", indent=0.55):
        for it in items:
            p = self.d.add_paragraph(); f = p.paragraph_format
            f.left_indent = Cm(indent); f.first_line_indent = Cm(-0.4); para_fmt(p, after=2)
            r = p.add_run(marker + "  "); r.font.color.rgb = rgb(RASP); r.bold = True
            add_runs(p, it, size=size)

    def checks(self, items, size=None):
        for it in items:
            p = self.d.add_paragraph(); f = p.paragraph_format
            f.left_indent = Cm(0.8); f.first_line_indent = Cm(-0.8); para_fmt(p, after=3)
            r = p.add_run("☐  "); r.font.name = "Segoe UI Symbol"; r.font.size = Pt(12); r.font.color.rgb = rgb(NAVY)
            add_runs(p, it, size=size)

    def spacer(self, pt=4):
        p = self.d.add_paragraph(); para_fmt(p, after=0, line=1.0); p.runs and None
        r = p.add_run(""); r.font.size = Pt(pt); p.paragraph_format.space_after = Pt(0)
        pPr = p._p.get_or_add_pPr(); sp = pPr.find(qn("w:spacing")); sp.set(qn("w:line"), str(int(pt * 20))); sp.set(qn("w:lineRule"), "exact")
        return p

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------- blocks
    def title_block(self, kicker, title, subtitle):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; shade(c, RASP); cell_margins(c, 200, 200, 280, 280); borders(c)
        p = c.paragraphs[0]; add_runs(p, kicker.upper(), size=9.5, color="FFD6E2", bold=True); para_fmt(p, after=2)
        p = c.add_paragraph(); add_runs(p, title, size=22, color="FFFFFF", bold=True); para_fmt(p, after=3, line=1.05)
        p = c.add_paragraph(); add_runs(p, subtitle, size=11, color="FFFFFF"); para_fmt(p, after=0)
        set_col_widths(t, [CONTENT_W]); self.spacer(6)

    def meta_table(self, rows):
        t = self.d.add_table(rows=len(rows), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (k, v) in enumerate(rows):
            a, b = t.rows[i].cells; no_split(t.rows[i])
            shade(a, "EEF1F4"); shade(b, "F8F9FB")
            for c in (a, b): cell_margins(c, 50, 50, 110, 110); borders(c, bottom=(4, "FFFFFF"))
            add_runs(a.paragraphs[0], k, size=9.5, color=NAVY, bold=True); para_fmt(a.paragraphs[0], after=0)
            add_runs(b.paragraphs[0], v, size=9.5); para_fmt(b.paragraphs[0], after=0)
        set_col_widths(t, [3.6, CONTENT_W - 3.6]); self.spacer(6)

    def callout(self, kind, title, lines, size=10):
        col = {"tip": GREEN, "warn": "C0392B", "note": NAVY, "book": PURPLE, "teach": AMBER}[kind]
        fill = {"tip": "E8F5EC", "warn": "FDECEA", "note": "E8EEF6", "book": "F0EAF6", "teach": "FFF4DB"}[kind]
        icon = {"tip": "TIP", "warn": "CAREFUL", "note": "NOTE", "book": "READ IN THE BOOK", "teach": "TEACHING TIP"}[kind]
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; no_split(t.rows[0]); shade(c, fill); cell_margins(c, 90, 90, 160, 140)
        borders(c, left=(36, col))
        p = c.paragraphs[0]; r = p.add_run(f"{icon}  "); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = rgb(col)
        if title: add_runs(p, title, size=size, bold=True, color=col)
        para_fmt(p, after=2, keep_next=True)
        first = True
        for ln in lines:
            if ln.startswith("- "):
                q = c.add_paragraph(); q.paragraph_format.left_indent = Cm(0.45); q.paragraph_format.first_line_indent = Cm(-0.35)
                rr = q.add_run("•  "); rr.font.color.rgb = rgb(col); rr.bold = True; add_runs(q, ln[2:], size=size); para_fmt(q, after=1.5, line=1.1)
            else:
                q = c.add_paragraph(); add_runs(q, ln, size=size); para_fmt(q, after=2, line=1.1)
        set_col_widths(t, [CONTENT_W]); self.spacer(5)

    def code(self, lines, caption=None, kind="cmd"):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; no_split(t.rows[0]); shade(c, "10202D" if kind == "term" else "F1F4F7"); cell_margins(c, 80, 80, 170, 120)
        borders(c, left=(24, "7CE0A3" if kind == "term" else NAVY))
        first = True
        for ln in lines:
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            if ln.startswith("# "):
                r = p.add_run(ln); r.font.name = MONO; r.font.size = Pt(9); r.font.color.rgb = rgb("7C8B99" if kind == "term" else "6B7785"); r.italic = True
            else:
                r = p.add_run(ln); r.font.name = MONO; r.font.size = Pt(9.6); r.bold = True if kind != "term" else False
                r.font.color.rgb = rgb("E7EEF5" if kind == "term" else "0F2742")
            r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
            para_fmt(p, after=0, line=1.15)
        set_col_widths(t, [CONTENT_W]); self.spacer(5)

    def step(self, n, title, body_lines=(), color=RASP):
        t = self.d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        a, b = t.rows[0].cells; no_split(t.rows[0])
        shade(a, color); cell_margins(a, 60, 60, 60, 60); a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        borders(a); borders(b, bottom=(4, "D9DEE4"))
        cell_margins(b, 60, 60, 150, 80)
        p = a.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, str(n), size=15, color="FFFFFF", bold=True); para_fmt(p, after=0)
        p = b.paragraphs[0]; add_runs(p, title, size=11, bold=True, color=NAVY); para_fmt(p, after=2, keep_next=True)
        for ln in body_lines:
            if ln.startswith("- "):
                q = b.add_paragraph(); q.paragraph_format.left_indent = Cm(0.45); q.paragraph_format.first_line_indent = Cm(-0.35)
                rr = q.add_run("•  "); rr.font.color.rgb = rgb(color); rr.bold = True; add_runs(q, ln[2:]); para_fmt(q, after=1.5, line=1.1)
            else:
                q = b.add_paragraph(); add_runs(q, ln); para_fmt(q, after=2, line=1.1)
        set_col_widths(t, [1.0, CONTENT_W - 1.0]); self.spacer(3)

    def table(self, header, rows, widths, size=9.5, header_fill=NAVY, zebra=True):
        t = self.d.add_table(rows=1 + len(rows), cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; shade(c, header_fill); cell_margins(c, 60, 60, 100, 100); borders(c)
            add_runs(c.paragraphs[0], h, size=size, color="FFFFFF", bold=True); para_fmt(c.paragraphs[0], after=0)
        repeat_header(t.rows[0])
        for i, row in enumerate(rows):
            no_split(t.rows[i + 1])
            for j, v in enumerate(row):
                c = t.rows[i + 1].cells[j]; shade(c, "F5F7FA" if (zebra and i % 2) else "FFFFFF"); cell_margins(c, 55, 55, 100, 100)
                borders(c, bottom=(4, "D9DEE4"))
                lines = v.split("\n")
                for k, ln in enumerate(lines):
                    p = c.paragraphs[0] if k == 0 else c.add_paragraph(); add_runs(p, ln, size=size); para_fmt(p, after=0, line=1.1)
        set_col_widths(t, widths); self.spacer(5)

    def figure(self, path, width_cm, caption, source=None, number=True):
        w_px, h_px = Image.open(path).size
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; para_fmt(p, before=3, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(path, width=Cm(width_cm))
        if number: self.fig_n += 1
        cap = self.d.add_paragraph(); para_fmt(cap, after=1 if source else 7, align=WD_ALIGN_PARAGRAPH.CENTER, keep_lines=True)
        r = cap.add_run((f"Figure {self.fig_n}.  " if number else "")); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = rgb(NAVY)
        add_runs(cap, caption, size=9, color=INK)
        if source:
            sp = self.d.add_paragraph(); para_fmt(sp, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_runs(sp, source, size=8, color=GREY, italic=True)
        dpi = w_px / (width_cm / 2.54)
        return dpi

    def figure_pair(self, left, right, width_cm, cap_left, cap_right, source_left=None):
        """Two figures side by side in a borderless table; each = (path)."""
        t = self.d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_split(t.rows[0])
        for c, path, cap, src in zip(t.rows[0].cells, (left, right), (cap_left, cap_right), (source_left, None)):
            cell_margins(c, 40, 40, 60, 60); borders(c)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; para_fmt(p, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(path, width=Cm(width_cm))
            self.fig_n += 1
            q = c.add_paragraph(); para_fmt(q, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = q.add_run(f"Figure {self.fig_n}.  "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = rgb(NAVY)
            add_runs(q, cap, size=9)
            if src:
                z = c.add_paragraph(); para_fmt(z, after=0, align=WD_ALIGN_PARAGRAPH.CENTER); add_runs(z, src, size=7.5, color=GREY, italic=True)
        set_col_widths(t, [CONTENT_W / 2, CONTENT_W / 2]); self.spacer(5)

    def save(self, path): self.d.save(path)
