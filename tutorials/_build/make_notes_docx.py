"""Wrap the 300-dpi A4 page images of a sketch-note edition into an A4 .docx (one full-bleed image per page).
usage: make_notes_docx.py OUT.docx page1.png page2.png ..."""
import sys
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_BREAK
out, *pages = sys.argv[1:]
d = Document(); s = d.sections[0]
s.page_width = Cm(21.0); s.page_height = Cm(29.7)
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(0); s.header_distance = s.footer_distance = Cm(0)
for i, p in enumerate(pages):
    para = d.add_paragraph()
    pf = para.paragraph_format; pf.space_before = pf.space_after = 0; pf.line_spacing = 1.0
    para.add_run().add_picture(p, width=Cm(21.0), height=Cm(29.7))
    if i < len(pages) - 1: para.add_run().add_break(WD_BREAK.PAGE)
d.save(out)
